"""Generate a deliberately non-compliant DOCX for the red path test.

Each anti-pattern below is annotated with the rule it should trigger
under the official KM typology, so we can certify that the audit
catches it and that ``patch_document`` repairs it whenever possible.
"""
from __future__ import annotations

import os

from docx import Document


BAD_FILENAME = "guidelines.docx"  # A1 KO : pas de date, pas de format normalisé


def build_bad() -> Document:
    doc = Document()

    # Aucun cartouche structuré           → A2 KO
    # Pas de section "Objectif"           → A3 KO
    # Pas de "Glossaire"                  → A4 KO
    # Pas de styles Heading               → B5 KO et B6 KO
    # ✓ glissé dans une phrase            → C9 KO
    # KM et IA non développés             → C7 KO
    doc.add_paragraph(
        "Notes diverses sur le projet et statut courant ✓ avec quelques "
        "détails supplémentaires sur les sigles KM et IA non développés."
    )

    # H28 KO : email non-auteur, IBAN, téléphone, adresse postale.
    # Aucun label « Auteur » présent → l'email n'est pas blanchi par H28.
    doc.add_paragraph(
        "Référent externe : alice.client@externe.example. "
        "IBAN fournisseur : FR7612345678901234567890123. "
        "Téléphone direct : +33 6 12 34 56 78. "
        "Adresse : 12 rue Lafayette."
    )

    # G26 KO : URL avec paramètres de tracking
    doc.add_paragraph(
        "Voir https://example.com/page?utm_source=foo&token=secret&id=42 pour plus."
    )

    return doc


def main() -> str:
    doc = build_bad()
    out_dir = os.path.dirname(os.path.abspath(__file__))
    out_path = os.path.join(out_dir, BAD_FILENAME)
    doc.save(out_path)
    print(f"Generated: {out_path}")
    return out_path


if __name__ == "__main__":
    main()
