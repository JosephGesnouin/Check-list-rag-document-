"""Remediation engine.

Given a DocumentAuditResult, build a ``RemediationPlan`` of concrete
actions. Then either:
  * apply the plan in-place to a DOCX file (the engine has direct write
    capability), or
  * render a Markdown remediation guide (works for any format).

The plan is intentionally explicit and idempotent: the user picks which
actions to apply via ``ActionKind`` toggles.
"""
from __future__ import annotations

import datetime as _dt
import io
import os
import re
from dataclasses import dataclass, field
from enum import Enum
from typing import Any, Dict, List, Optional

from .config import (
    EMOJI_OR_SYMBOL_RE,
    NAMING_REGEX,
    REQUIRED_CARTOUCHE_FIELDS,
    SYMBOL_REPLACEMENTS,
    URL_RE,
)
from .models import DocumentAuditResult, RuleResult, Status
from .sensitive import detect_sensitive, mask_email, mask_iban, mask_phone

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover
    DocxDocument = None


# ---------------------------------------------------------------------------
# Plan model
# ---------------------------------------------------------------------------


class ActionKind(str, Enum):
    RENAME_FILE = "rename_file"
    INJECT_CARTOUCHE = "inject_cartouche"
    INJECT_OBJECTIVE = "inject_objective"
    INJECT_GLOSSARY = "inject_glossary"
    REPLACE_SYMBOLS = "replace_symbols"
    CLEAN_URLS = "clean_urls"
    REDACT_SENSITIVE = "redact_sensitive"


@dataclass
class RemediationAction:
    kind: ActionKind
    rule_id: str
    title: str
    description: str
    payload: Dict[str, Any] = field(default_factory=dict)
    applicable_formats: tuple = ("docx",)  # what patch_document supports

    def supports_patch(self, file_type: str) -> bool:
        return file_type in self.applicable_formats


@dataclass
class RemediationPlan:
    file_name: str
    file_type: str
    actions: List[RemediationAction] = field(default_factory=list)

    def by_rule(self) -> Dict[str, RemediationAction]:
        return {a.rule_id: a for a in self.actions}


# ---------------------------------------------------------------------------
# Plan builder
# ---------------------------------------------------------------------------


def build_plan(result: DocumentAuditResult) -> RemediationPlan:
    """Translate failed/remediable rules into a concrete action list."""
    plan = RemediationPlan(file_name=result.file_name, file_type=result.file_type)
    failing = {r.rule_id: r for r in result.rules if r.status in (Status.FAIL, Status.WARN, Status.NOT_VERIFIABLE)}

    if "A1" in failing and failing["A1"].status is Status.FAIL:
        plan.actions.append(_action_rename(result, failing["A1"]))
    if "A2" in failing:
        plan.actions.append(_action_cartouche(result, failing["A2"]))
    if "A3" in failing and failing["A3"].status is Status.FAIL:
        plan.actions.append(_action_objective(failing["A3"]))
    if "A5" in failing and failing["A5"].status is Status.FAIL:
        plan.actions.append(_action_glossary(failing["A5"]))
    if any(rid in failing and failing[rid].status is Status.FAIL for rid in ("B9", "D16")):
        plan.actions.append(_action_replace_symbols())
    if "F26" in failing and failing["F26"].status is Status.FAIL:
        plan.actions.append(_action_clean_urls())
    if "H34" in failing and failing["H34"].status is Status.FAIL:
        plan.actions.append(_action_redact_sensitive())

    return plan


def _action_rename(result: DocumentAuditResult, rule: RuleResult) -> RemediationAction:
    suggested = _suggest_filename(result.file_name, result.file_type)
    return RemediationAction(
        kind=ActionKind.RENAME_FILE,
        rule_id=rule.rule_id,
        title="Renommer le fichier au format normalisé",
        description=f"Renommer en {suggested}",
        payload={"new_name": suggested},
        applicable_formats=("docx", "pptx", "xlsx", "pdf"),
    )


def _action_cartouche(result: DocumentAuditResult, rule: RuleResult) -> RemediationAction:
    return RemediationAction(
        kind=ActionKind.INJECT_CARTOUCHE,
        rule_id=rule.rule_id,
        title="Insérer un cartouche pré-rempli",
        description="Ajoute en tête de document une section 'Cartouche' avec les champs obligatoires à compléter.",
        payload={"fields": list(REQUIRED_CARTOUCHE_FIELDS)},
        applicable_formats=("docx",),
    )


def _action_objective(rule: RuleResult) -> RemediationAction:
    return RemediationAction(
        kind=ActionKind.INJECT_OBJECTIVE,
        rule_id=rule.rule_id,
        title="Insérer une section Objectif",
        description="Ajoute un paragraphe 'Objectif' à compléter en début de document.",
        applicable_formats=("docx",),
    )


def _action_glossary(rule: RuleResult) -> RemediationAction:
    return RemediationAction(
        kind=ActionKind.INJECT_GLOSSARY,
        rule_id=rule.rule_id,
        title="Insérer un Glossaire en fin de document",
        description="Ajoute une section Glossaire (titre + tableau acronyme/définition vide).",
        applicable_formats=("docx",),
    )


def _action_replace_symbols() -> RemediationAction:
    return RemediationAction(
        kind=ActionKind.REPLACE_SYMBOLS,
        rule_id="B9/D16",
        title="Remplacer les symboles (✓, ✗, ➜, ...) par des mots",
        description="Substitue les pictogrammes par leur équivalent textuel (oui/non/->) dans paragraphes et tableaux.",
        applicable_formats=("docx",),
    )


def _action_clean_urls() -> RemediationAction:
    return RemediationAction(
        kind=ActionKind.CLEAN_URLS,
        rule_id="F26",
        title="Nettoyer les paramètres de tracking dans les URLs",
        description="Retire les paramètres utm_/token/session/gclid/fbclid/auth/tracking des URLs en clair et des hyperliens.",
        applicable_formats=("docx",),
    )


def _action_redact_sensitive() -> RemediationAction:
    return RemediationAction(
        kind=ActionKind.REDACT_SENSITIVE,
        rule_id="H34",
        title="Masquer les données sensibles détectées",
        description="Remplace emails/téléphones/IBAN par leur version masquée (j***@domaine.com, etc.).",
        applicable_formats=("docx",),
    )


def _suggest_filename(file_name: str, file_type: str) -> str:
    if NAMING_REGEX.match(file_name):
        return file_name
    base = os.path.splitext(os.path.basename(file_name))[0]
    safe = re.sub(r"[^A-Za-z0-9À-ÿ\s-]", "", base).strip() or "document"
    parts = safe.split()
    subject = (parts[0] if parts else "Sujet")[:40]
    extra = "_".join(parts[1:3]) if len(parts) > 1 else ""
    today = _dt.date.today().strftime("%Y%m%d")
    pieces = [today, subject, "KMDoc"]
    if extra:
        pieces.append(extra)
    return "_".join(pieces) + "." + (file_type or "docx")


# ---------------------------------------------------------------------------
# Apply plan to DOCX
# ---------------------------------------------------------------------------


@dataclass
class PatchOutcome:
    file_name: str
    patched_bytes: Optional[bytes]
    applied: List[str]
    skipped: List[str]
    error: Optional[str] = None


def patch_document(
    file_name: str,
    raw: bytes,
    plan: RemediationPlan,
    enabled_kinds: Optional[List[ActionKind]] = None,
) -> PatchOutcome:
    """Apply selected actions to a DOCX. Returns patched bytes + applied list.

    For non-DOCX formats, no in-place patching is performed (use
    ``render_markdown_guide`` instead). The rename action is reported via
    ``patched_bytes`` keeping the same content but with the suggested name.
    """
    enabled = set(enabled_kinds or [a.kind for a in plan.actions])
    file_type = (plan.file_type or "").lower()
    applied: List[str] = []
    skipped: List[str] = []

    new_name = file_name
    rename = next(
        (a for a in plan.actions if a.kind is ActionKind.RENAME_FILE and a.kind in enabled),
        None,
    )
    if rename is not None:
        new_name = rename.payload.get("new_name", file_name)
        applied.append(rename.title)

    if file_type != "docx" or DocxDocument is None:
        # In-place patching not supported for this format.
        for a in plan.actions:
            if a.kind in enabled and a.kind is not ActionKind.RENAME_FILE:
                skipped.append(a.title)
        return PatchOutcome(
            file_name=new_name,
            patched_bytes=raw if rename else None,
            applied=applied,
            skipped=skipped,
            error=None if file_type == "docx" else (
                f"Patch in-place non supporté pour .{file_type}. "
                "Voir le guide Markdown pour la procédure manuelle."
            ),
        )

    try:
        doc = DocxDocument(io.BytesIO(raw))
        for action in plan.actions:
            if action.kind not in enabled:
                continue
            if action.kind is ActionKind.RENAME_FILE:
                continue
            if action.kind is ActionKind.INJECT_CARTOUCHE:
                _docx_inject_cartouche(doc, action.payload.get("fields", REQUIRED_CARTOUCHE_FIELDS))
            elif action.kind is ActionKind.INJECT_OBJECTIVE:
                _docx_inject_objective(doc)
            elif action.kind is ActionKind.INJECT_GLOSSARY:
                _docx_inject_glossary(doc)
            elif action.kind is ActionKind.REPLACE_SYMBOLS:
                _docx_replace_symbols(doc)
            elif action.kind is ActionKind.CLEAN_URLS:
                _docx_clean_urls(doc)
            elif action.kind is ActionKind.REDACT_SENSITIVE:
                _docx_redact_sensitive(doc)
            else:
                skipped.append(action.title)
                continue
            applied.append(action.title)

        out = io.BytesIO()
        doc.save(out)
        return PatchOutcome(
            file_name=new_name,
            patched_bytes=out.getvalue(),
            applied=applied,
            skipped=skipped,
        )
    except Exception as exc:  # pragma: no cover - defensive
        return PatchOutcome(
            file_name=new_name,
            patched_bytes=None,
            applied=applied,
            skipped=skipped,
            error=f"Erreur lors du patch: {exc}",
        )


# ---------------------------------------------------------------------------
# DOCX patch helpers
# ---------------------------------------------------------------------------


def _docx_paragraphs_iter(doc):
    yield from doc.paragraphs
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _set_paragraph_text(para, new_text: str) -> None:
    if not para.runs:
        para.add_run(new_text)
        return
    para.runs[0].text = new_text
    for run in para.runs[1:]:
        run.text = ""


def _docx_replace_symbols(doc) -> None:
    pattern = re.compile("|".join(re.escape(k) for k in SYMBOL_REPLACEMENTS))

    def _sub(text: str) -> str:
        new = pattern.sub(lambda m: SYMBOL_REPLACEMENTS[m.group(0)], text)
        # Sweep any remaining decorative symbols (not in dict) – drop them.
        return EMOJI_OR_SYMBOL_RE.sub("", new)

    for para in _docx_paragraphs_iter(doc):
        if not para.text:
            continue
        new = _sub(para.text)
        if new != para.text:
            _set_paragraph_text(para, new)


def _docx_clean_urls(doc) -> None:
    from .rules.urls import clean_url

    for para in _docx_paragraphs_iter(doc):
        if URL_RE.search(para.text or ""):
            new = URL_RE.sub(lambda m: clean_url(m.group(0))[0], para.text)
            if new != para.text:
                _set_paragraph_text(para, new)

    # Also clean hyperlink targets.
    rels = doc.part.rels
    for rel in rels.values():
        if rel.reltype.endswith("/hyperlink") and rel.target_ref and rel.target_ref.startswith("http"):
            cleaned, _ = clean_url(rel.target_ref)
            if cleaned != rel.target_ref:
                rel._target = cleaned  # noqa: SLF001 - python-docx exposes _target


def _docx_redact_sensitive(doc) -> None:
    full_text_blocks = [p.text for p in _docx_paragraphs_iter(doc)]
    findings = detect_sensitive("\n".join(full_text_blocks))
    if not any(findings.values()):
        return
    masks = {}
    for em in findings.get("emails", []):
        masks[em] = mask_email(em)
    for ph in findings.get("phones", []):
        masks[ph] = mask_phone(ph)
    for ib in findings.get("ibans", []):
        masks[ib] = mask_iban(ib)

    for para in _docx_paragraphs_iter(doc):
        text = para.text or ""
        if not text:
            continue
        new = text
        for raw, masked in masks.items():
            if raw in new:
                new = new.replace(raw, masked)
        if new != text:
            _set_paragraph_text(para, new)


def _docx_inject_cartouche(doc, fields) -> None:
    if _has_marker(doc, "[CARTOUCHE]"):
        return
    body = doc.element.body
    cartouche_paras = []
    cartouche_paras.append(_make_heading(doc, "Cartouche de présentation du document", level=1))
    cartouche_paras.append(_make_paragraph(doc, "[CARTOUCHE] (généré automatiquement – à compléter)"))
    for field_name in fields:
        cartouche_paras.append(_make_paragraph(doc, f"{field_name}* : "))
    cartouche_paras.append(_make_paragraph(doc, ""))
    # Insert at the very top
    for para in reversed(cartouche_paras):
        body.insert(0, para._element)


def _docx_inject_objective(doc) -> None:
    if _has_marker(doc, "[OBJECTIF]"):
        return
    body = doc.element.body
    paras = [
        _make_heading(doc, "Objectif", level=2),
        _make_paragraph(doc, "[OBJECTIF] Décrire ici le but du document, le public cible et le périmètre."),
        _make_paragraph(doc, ""),
    ]
    # Insert after a likely existing cartouche (top of doc)
    insert_at = 0
    for idx, child in enumerate(list(body)[:10]):
        if child.tag.endswith("p"):
            text = "".join(child.itertext()) or ""
            if text.strip():
                insert_at = idx + 1
    for i, para in enumerate(paras):
        body.insert(insert_at + i, para._element)


def _docx_inject_glossary(doc) -> None:
    if _has_marker(doc, "[GLOSSAIRE]"):
        return
    doc.add_paragraph()
    doc.add_heading("Glossaire", level=1)
    doc.add_paragraph("[GLOSSAIRE] Lister ci-dessous les acronymes et termes techniques.")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Acronyme / Terme"
    hdr[1].text = "Définition"
    sample = table.rows[1].cells
    sample[0].text = "..."
    sample[1].text = "..."


def _has_marker(doc, marker: str) -> bool:
    for para in doc.paragraphs:
        if marker in (para.text or ""):
            return True
    return False


def _make_paragraph(doc, text: str):
    para = doc.add_paragraph(text)
    # Move it out of body tail; we'll reinsert manually
    doc.element.body.remove(para._element)
    return para


def _make_heading(doc, text: str, level: int = 1):
    para = doc.add_heading(text, level=level)
    doc.element.body.remove(para._element)
    return para


# ---------------------------------------------------------------------------
# Markdown remediation guide (works for any format)
# ---------------------------------------------------------------------------


def render_markdown_guide(result: DocumentAuditResult, plan: RemediationPlan) -> str:
    lines: List[str] = []
    lines.append(f"# Guide de remédiation – {result.file_name}")
    lines.append("")
    lines.append(f"- **Verdict actuel**: {result.verdict}")
    lines.append(f"- **Score**: {result.score}/100")
    lines.append(f"- **Date d'audit**: {result.audit_date}")
    lines.append("")

    if not plan.actions:
        lines.append("Aucune action corrective requise.")
        return "\n".join(lines)

    lines.append("## Actions recommandées")
    lines.append("")
    for action in plan.actions:
        applicable = "oui" if action.supports_patch(plan.file_type) else "non (manuel)"
        lines.append(f"### {action.title}")
        lines.append(f"- Règle: `{action.rule_id}`")
        lines.append(f"- Type d'action: `{action.kind.value}`")
        lines.append(f"- Patch automatique sur ce format: **{applicable}**")
        lines.append(f"- Description: {action.description}")
        if action.payload:
            lines.append(f"- Détails: `{action.payload}`")
        lines.append("")

    lines.append("## Points résiduels (vérification manuelle)")
    lines.append("")
    for rule in result.rules:
        if rule.status in (Status.FAIL, Status.WARN, Status.NOT_VERIFIABLE):
            if rule.rule_id in plan.by_rule():
                continue
            lines.append(f"- **{rule.rule_id} {rule.title}** ({rule.status.value}) – {rule.recommendation or rule.evidence}")
    return "\n".join(lines)
